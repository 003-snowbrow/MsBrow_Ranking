import sys, json
import openpyxl

src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, data_only=True)
target = None
for name in wb.sheetnames:
    if name.strip() == "一部考核總表":
        target = wb[name]
        break
if target is None:
    print("SHEETS:", wb.sheetnames)
    raise SystemExit("sheet 一部考核總表 not found")

rows = []
for r in range(1, 41):
    row = []
    for c in range(1, 20):  # A..S
        v = target.cell(row=r, column=c).value
        row.append(v)
    rows.append(row)

date_range = None
for name in wb.sheetnames:
    ws = wb[name]
    if ws.cell(row=1, column=1).value and "员工业绩统计表" in str(ws.cell(row=1, column=1).value):
        date_range = ws.cell(row=3, column=5).value
        break

json.dump({"sheets": wb.sheetnames, "date_range": date_range, "rows": rows}, open(out, "w"), ensure_ascii=False, indent=1, default=str)
print("OK", out, "date_range:", date_range)
